"""Rebuild reference list in APA 7 (full author lists up to 20, italics), fix in-text disambiguation."""
import docx, copy, re, json
SRC = 'research/tmp/user_latest.docx'; OUT = 'research/tmp/apa.docx'
d = docx.Document(SRC); P = d.paragraphs

def initials(full):
    parts = full.replace('.', '. ').split()
    *given, last = parts
    if full == 'Michael R. Lyu': pass
    ini = ' '.join(('-'.join(x[0] + '.' for x in g.split('-')) if not g.endswith('.') else g) for g in given)
    return last, ini

def fmt_authors(names):
    out = []
    for n in names:
        # handle multi-word surnames
        special = {'Manuel Del Verme': ('Del Verme', 'M.'), 'José Ramón Enríquez': ('Enríquez', 'J. R.'),
                   'Benjamin Van Durme': ('Van Durme', 'B.'), "Kenton O'Hara": ("O'Hara", 'K.'),
                   'Gorkem Ozer Yilmaz': ('Ozer Yilmaz', 'G.'), 'Hao Yang Lu': ('Lu', 'H. Y.')}
        last, ini = special.get(n) or initials(n)
        out.append(f'{last}, {ini}')
    if len(out) == 1: return out[0]
    if len(out) <= 20: return ', '.join(out[:-1]) + ', & ' + out[-1]
    return ', '.join(out[:19]) + ', . . . ' + out[-1]

A = json.load(open('research/tmp/authors.json', encoding='utf-8'))
# (authors, year, [(text, italic)...])  — italics per APA 7
R = [
 (A['Alves'], 2024, [('Cost-sensitive learning to defer to multiple experts with workload constraints', 1), (' (arXiv:2403.06906). arXiv. https://doi.org/10.48550/arXiv.2403.06906', 0)]),
 (['Michelle Brachman', 'Amina El-Ashry', 'Casey Dugan', 'Werner Geyer'], 2025, [('Current and future use of large language models for knowledge work. ', 0), ('Proceedings of the ACM on Human-Computer Interaction, 9', 1), ('(7). https://doi.org/10.1145/3757403', 0)]),
 (['Yulin Chen', 'Yang Zhan', 'Zhuoran Lu', 'Qiao Jin'], 2026, [('When single-user-oriented LLM-based assistants involve others: A scoping review of pathways, risks, and responses', 1), (' (arXiv:2609.14062). arXiv. https://doi.org/10.48550/arXiv.2609.14062', 0)]),
 (['Hyesun Chung', 'Timothy Holder', 'Julie Shah', 'X. Jessie Yang'], 2024, [('Developing a team classification scheme for human-agent teaming. ', 0), ('Proceedings of the Human Factors and Ergonomics Society Annual Meeting, 68', 1), ('(1), 1394–1399. https://doi.org/10.1177/10711813241260387', 0)]),
 (A['Counts'], 2026, [('AI in the enterprise: How people use M365 Copilot Chat', 1), (' (arXiv:2605.23958). arXiv. https://doi.org/10.48550/arXiv.2605.23958', 0)]),
 (['Ruomeng Ding', 'Tianwei Gao', 'Thomas P. Zollo', 'Eitan Bachmat', 'Richard Zemel', 'Zhun Deng'], 2026, [('Whom to query for what: Adaptive group elicitation via multi-turn LLM interactions. In ', 0), ('Proceedings of the 43rd International Conference on Machine Learning', 1), ('. https://doi.org/10.48550/arXiv.2602.14279', 0)]),
 (A['Drouin'], 2024, [('WorkArena: How capable are web agents at solving common knowledge work tasks? In ', 0), ('Proceedings of the 41st International Conference on Machine Learning', 1), ('. https://doi.org/10.48550/arXiv.2403.07718', 0)]),
 (['Wangxuan Fan', 'Xiaoyu Nie', 'Zhongxiang Dai'], 2026, [('Harness-MU: A safe, governed, and effective harness for multi-user LLM agents', 1), (' (arXiv:2606.21856). arXiv. https://doi.org/10.48550/arXiv.2606.21856', 0)]),
 (A['Feng'], 2024, [('Large language model-based human-agent collaboration for complex task solving. In ', 0), ('Findings of the Association for Computational Linguistics: EMNLP 2024', 1), (' (pp. 1336–1357). Association for Computational Linguistics. https://doi.org/10.18653/v1/2024.findings-emnlp.72', 0)]),
 (A['Gao'], 2024, [('Aligning LLM agents by learning latent preference from user edits. In ', 0), ('Advances in Neural Information Processing Systems 37', 1), ('. https://doi.org/10.48550/arXiv.2404.15269', 0)]),
 (A['Huang'], 2025, [('CRMArena: Understanding the capacity of LLM agents to perform professional CRM tasks in realistic environments. In ', 0), ('Proceedings of the 2025 Conference of the Nations of the Americas Chapter of the Association for Computational Linguistics', 1), ('. Association for Computational Linguistics. https://aclanthology.org/2025.naacl-long.194/', 0)]),
 (['Harsh Jhamtani', 'Jacob Andreas', 'Benjamin Van Durme'], 2025, [('LLM agents for coordinating multi-user information gathering. In ', 0), ('Findings of the Association for Computational Linguistics: ACL 2025', 1), ('. Association for Computational Linguistics. https://doi.org/10.18653/v1/2025.findings-acl.916', 0)]),
 (A['Lu'], 2025, [('Proactive agent: Shifting LLM agents from reactive responses to active assistance. In ', 0), ('Proceedings of the 13th International Conference on Learning Representations', 1), ('. https://doi.org/10.48550/arXiv.2410.12361', 0)]),
 (['Manqing Mao', 'Paishun Ting', 'Yijian Xiang', 'Mingyang Xu', 'Julia Chen', 'Jianzhe Lin'], 2024, [('Multi-user chat assistant (MUCA): A framework using LLMs to facilitate group conversations', 1), (' (arXiv:2401.04883). arXiv. https://doi.org/10.48550/arXiv.2401.04883', 0)]),
 (['Yannis Montreuil', 'Axel Carlier', 'Lai Xing Ng', 'Wei Tsang Ooi'], 2026, [('Why ask one when you can ask k? Learning-to-defer to the top-k experts. In ', 0), ('Proceedings of the 14th International Conference on Learning Representations', 1), ('. https://doi.org/10.48550/arXiv.2504.12988', 0)]),
 (["Kenton O'Hara", 'Michael Massimi', 'Miftah Khan'], 2026, [('Exploring how knowledge workers situate an organisationally embedded LLM agent in their collaborative work practices. ', 0), ('Computer Supported Cooperative Work (CSCW)', 1), ('. https://doi.org/10.1007/s10606-026-09549-1', 0)]),
 (A['Ren2023'], 2023, [('Robots that ask for help: Uncertainty alignment for large language model planners. In ', 0), ('Proceedings of the 7th Conference on Robot Learning', 1), ('. https://doi.org/10.48550/arXiv.2307.01928', 0)]),
 (A['Ren2026'], 2026, [('GateMem: Benchmarking memory governance in multi-principal shared-memory agents', 1), (' (arXiv:2606.18829). arXiv. https://doi.org/10.48550/arXiv.2606.18829', 0)]),
 (['Alireza Rezazadeh', 'Zichao Li', 'Ange Lou', 'Yuying Zhao', 'Wei Wei', 'Yujia Bao'], 2025, [('Collaborative memory: Multi-user memory sharing in LLM agents with dynamic access control', 1), (' (arXiv:2505.18279). arXiv. https://doi.org/10.48550/arXiv.2505.18279', 0)]),
 (['Debdeep Sanyal', 'Umakanta Maharana', 'Yash Sinha', 'Hong Ming Tan', 'Shirish Karande', 'Mohan Kankanhalli'], 2025, [('OrgAccess: A benchmark for role based access control in organization scale LLMs', 1), (' (arXiv:2505.19165). arXiv. https://doi.org/10.48550/arXiv.2505.19165', 0)]),
 (['Yijia Shao', 'Vinay Samuel', 'Yucheng Jiang', 'John Yang', 'Diyi Yang'], 2024, [('Collaborative Gym: A framework for enabling and evaluating human-agent collaboration', 1), (' (arXiv:2412.15701). arXiv. https://doi.org/10.48550/arXiv.2412.15701', 0)]),
 (A['Suri'], 2025, [('Structured uncertainty guided clarification for LLM agents', 1), (' (arXiv:2511.08798). arXiv. https://doi.org/10.48550/arXiv.2511.08798', 0)]),
 (['Harsh Vishwakarma', 'Ankush Agarwal', 'Ojas Patil', 'Chaitanya Devaguptapu', 'Mahesh Chandran'], 2025, [('Can LLMs help you at work? A sandbox for evaluating LLM agents in enterprise environments. In ', 0), ('Proceedings of the 2025 Conference on Empirical Methods in Natural Language Processing', 1), ('. Association for Computational Linguistics. https://doi.org/10.48550/arXiv.2510.27287', 0)]),
 (A['Wang2025'], 2025, [('Learning to ask: When LLM agents meet unclear instruction. In ', 0), ('Proceedings of the 2025 Conference on Empirical Methods in Natural Language Processing', 1), ('. Association for Computational Linguistics. https://aclanthology.org/2025.emnlp-main.1104/', 0)]),
 (['Yihe Wang', 'Kehua Lei', 'Sheng-Yang Chiu', 'Katherine Isbister', 'David T. Lee', 'Kathryn E. Ringland'], 2026, [('SeeSawBot: An LLM-driven chatbot mediating across private and shared Slack channels to support team dynamics. In ', 0), ('Proceedings of the 2026 CHI Conference on Human Factors in Computing Systems', 1), ('. ACM. https://doi.org/10.1145/3772318.3791880', 0)]),
 (['Yaozu Wu', 'Wei-Xing Huang', 'Jizhou Guo', 'Dongyuan Li', 'Renhe Jiang'], 2026, [('HAS-Bench: Evaluating LLM-based human-agent systems under configurable human participation', 1), (' (arXiv:2607.04329). arXiv. https://doi.org/10.48550/arXiv.2607.04329', 0)]),
 (A['Xu2025'], 2025, [('TheAgentCompany: Benchmarking LLM agents on consequential real world tasks. In ', 0), ('Advances in Neural Information Processing Systems 38: Datasets and Benchmarks Track', 1), ('. https://doi.org/10.48550/arXiv.2412.14161', 0)]),
 (A['Xu2026'], 2026, [('Toward personalized LLM-powered agents: Foundations, evaluation, and future directions', 1), (' (arXiv:2602.22680). arXiv. https://doi.org/10.48550/arXiv.2602.22680', 0)]),
 (['Shu Yang', 'Difei Xu', 'Jiaxin Pei', 'Di Wang'], 2026, [('ProACT: Towards breakdown-aware proactive agent in multi-user collaboration', 1), (' (arXiv:2607.03730). arXiv. https://doi.org/10.48550/arXiv.2607.03730', 0)]),
 (A['Yang2026'], 2026, [('Multi-user large language model agents', 1), (' (arXiv:2604.08567). arXiv. https://doi.org/10.48550/arXiv.2604.08567', 0)]),
 (A['Zollo'], 2025, [('PersonalLLM: Tailoring LLMs to individual preferences. In ', 0), ('Proceedings of the 13th International Conference on Learning Representations', 1), ('. https://doi.org/10.48550/arXiv.2409.20296', 0)]),
 (A['Zou'], 2025, [('A call for collaborative intelligence: Why human-agent systems should precede AI autonomy', 1), (' (arXiv:2506.09420). arXiv. https://doi.org/10.48550/arXiv.2506.09420', 0)]),
]

# in-text fixes (APA 7 disambiguation)
def rep_all(old, new):
    n = 0
    for p in d.paragraphs:
        for r in p.runs:
            if old in r.text: r.text = r.text.replace(old, new); n += 1
    return n
n1 = rep_all('(Yang et al., 2026)', '(Yang, Zhu, et al., 2026)')
n2 = rep_all('Yang et al., 2026;', 'Yang, Zhu, et al., 2026;')
n3 = rep_all('; Yang et al., 2026', '; Yang, Zhu, et al., 2026')
n4 = rep_all('Xu, Chen, et al., 2026', 'Xu et al., 2026')
print('intext fixes', n1, n2, n3, n4)

P = d.paragraphs
k = [i for i, p in enumerate(P) if p.text.strip() == 'References'][0]
old = [p for p in P[k + 1:]]
tmpl = old[1]._p if len(old) > 1 else old[0]._p
for p in old[1:]: p._p.getparent().remove(p._p)
cur = old[0]
for i, (auth, yr, parts) in enumerate(R):
    if i > 0:
        new = copy.deepcopy(tmpl); cur._p.addnext(new); cur = docx.text.paragraph.Paragraph(new, cur._parent)
    for r in list(cur.runs): r._r.getparent().remove(r._r)
    cur.add_run(f'{fmt_authors(auth)} ({yr}). ')
    for t, it in parts:
        r = cur.add_run(t); r.italic = bool(it)
d.save(OUT)
D = docx.Document(OUT); paras = [p.text for p in D.paragraphs]; k = paras.index('References')
print(len([p for p in paras[k + 1:] if p.strip()]), 'refs'); print(paras[k + 2][:200]); print(paras[k + 27][:300])
