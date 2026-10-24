import docx, copy, re
d = docx.Document('research/tmp/lr.docx'); P = d.paragraphs
refs = open('research/tmp/refs.txt', encoding='utf-8').read().strip().split('\n')
for p in P:
    for r in p.runs:
        if 'Zou et al., 2025a' in r.text: r.text = r.text.replace('Zou et al., 2025a', 'Zou et al., 2025')
idx = [i for i, p in enumerate(P) if p.text == 'References'][0]
old = P[idx + 1:]
tmpl = old[1]._p
for p in old[1:]: p._p.getparent().remove(p._p)
first = old[0]
for r in list(first.runs): r._r.getparent().remove(r._r)
first.add_run(refs[0]); last = first
for line in refs[1:]:
    new = copy.deepcopy(tmpl); last._p.addnext(new); q = docx.text.paragraph.Paragraph(new, last._parent)
    for r in list(q.runs): r._r.getparent().remove(r._r)
    q.add_run(line); last = q
d.save('research/tmp/lr2.docx')
D = docx.Document('research/tmp/lr2.docx'); paras = [p.text for p in D.paragraphs]
k = paras.index('References'); body = ' '.join(paras[:k]); reftext = '\n'.join(paras[k + 1:])
cits = set(re.findall(r"([A-Z][A-Za-z'’\-]+)(?: et al\.| & [A-Z][A-Za-z'’\-]+|, [A-Z][a-z]+, et al\.)?,? \(?(20\d\d)", body))
for a, y in sorted(cits):
    if not re.search(rf"^{re.escape(a)}[^\n]*?\({y}\)", reftext, re.M): print('MISSING REF?', a, y)
used = set(a for a, y in cits)
for line in refs:
    if line.split(',')[0] not in used: print('UNCITED?', line[:50])
print('refs', len(refs))
